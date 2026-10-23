#!/usr/bin/env python3
"""Structural comparison of the generated manuscript against the PRISM reference.

Reports side-by-side counts for sections, tables, figures, equations and body
size, plus an audit of citation placement against the rules fixed for this
paper:

    [1]-[5]    Introduction only
    [6]-[20]   Literature Survey prose, ascending order
    [21]-[30]  Literature Survey comparison table (with the Proposed row)

Usage:  python scripts/verify_paper.py [reference.docx] [generated.docx]
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from docx import Document

REF_DEFAULT = "PRISM - Copy.docx"
GEN_DEFAULT = "docs/SymptoGraph-LLM_Paper.docx"

HEADING_RE = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+(.*)$")
CITE_RE = re.compile(r"\[(\d+(?:\s*[-,]\s*\d+)*)\]")


def expand_citation(token: str) -> list[int]:
    """Turn the inside of a [...] into the integers it denotes."""
    out: list[int] = []
    for part in token.split(","):
        part = part.strip()
        if "-" in part:
            lo, hi = (int(x) for x in part.split("-", 1))
            out.extend(range(lo, hi + 1))
        elif part.isdigit():
            out.append(int(part))
    return out


def body_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def media_files(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as zf:
        return sorted(n for n in zf.namelist() if n.startswith("word/media/"))


def omml_count(path: Path) -> int:
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", "replace")
    return xml.count("<m:oMath")


def headings(doc: Document) -> list[tuple[str, str]]:
    found: list[tuple[str, str]] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text or len(text) > 90:
            continue
        m = HEADING_RE.match(text)
        if m and not text.endswith("."):
            found.append((m.group(1), m.group(2)))
    return found


def section_of(index: int, markers: dict[str, int]) -> str:
    """Which top-level section a paragraph index falls in."""
    current = "front"
    for name, start in sorted(markers.items(), key=lambda kv: kv[1]):
        if index >= start:
            current = name
    return current


def describe(label: str, path: Path) -> dict:
    doc = Document(str(path))
    text = body_text(doc)
    info = {
        "label": label,
        "paragraphs": len(doc.paragraphs),
        "chars": len(text),
        "tables": len(doc.tables),
        "table_shapes": [(len(t.rows), len(t.columns)) for t in doc.tables],
        "media": len(media_files(path)),
        "omml": omml_count(path),
        "headings": headings(doc),
        "fig_captions": [
            p.text.strip()
            for p in doc.paragraphs
            if re.match(r"^(Figure|Fig\.)\s*\d+", p.text.strip())
        ],
        "tbl_captions": [
            p.text.strip()
            for p in doc.paragraphs
            if re.match(r"^Table\s*\d+", p.text.strip())
        ],
        "placeholders": text.count("X.XXX"),
        "doc": doc,
        "text": text,
    }
    sec = doc.sections[0]
    info["page"] = (
        round(sec.page_width.inches, 2),
        round(sec.page_height.inches, 2),
    )
    info["margins"] = (
        round(sec.left_margin.inches, 2),
        round(sec.top_margin.inches, 2),
    )
    return info


def audit_citations(info: dict) -> list[str]:
    doc = info["doc"]
    problems: list[str] = []

    markers: dict[str, int] = {}
    for i, para in enumerate(doc.paragraphs):
        t = para.text.strip().lower()
        for key, pat in (
            ("intro", r"^1\.?\s+introduction"),
            ("survey", r"^2\.?\s+literature\s+survey"),
            ("method", r"^3\.?\s+proposed"),
            ("results", r"^4\.?\s+results"),
            ("concl", r"^5\.?\s+conclusion"),
            ("refs", r"^6\.?\s+references"),
        ):
            if key not in markers and re.match(pat, t):
                markers[key] = i

    per_section: dict[str, list[int]] = {}
    survey_order: list[int] = []
    for i, para in enumerate(doc.paragraphs):
        sec = section_of(i, markers)
        if sec == "refs":
            continue
        for m in CITE_RE.finditer(para.text):
            nums = expand_citation(m.group(1))
            per_section.setdefault(sec, []).extend(nums)
            if sec == "survey":
                survey_order.extend(nums)

    intro = sorted(set(per_section.get("intro", [])))
    survey_prose = sorted(set(survey_order))

    table_nums: list[int] = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for m in CITE_RE.finditer(cell.text):
                    table_nums.extend(expand_citation(m.group(1)))
    table_nums = sorted(set(table_nums))

    if intro != list(range(1, 6)):
        problems.append(f"introduction cites {intro}, expected [1..5]")
    if survey_prose != list(range(6, 21)):
        missing = sorted(set(range(6, 21)) - set(survey_prose))
        extra = sorted(set(survey_prose) - set(range(6, 21)))
        problems.append(
            f"survey prose cites {len(survey_prose)} refs; missing={missing} extra={extra}"
        )
    if survey_order != sorted(survey_order):
        first_bad = next(
            (
                (a, b)
                for a, b in zip(survey_order, survey_order[1:])
                if b < a
            ),
            None,
        )
        problems.append(f"survey prose not ascending, first descent at {first_bad}")
    if not set(range(21, 31)).issubset(set(table_nums)):
        problems.append(
            f"comparison table missing {sorted(set(range(21,31)) - set(table_nums))}"
        )

    ref_entries = 0
    if "refs" in markers:
        for para in doc.paragraphs[markers["refs"] + 1 :]:
            if CITE_RE.match(para.text.strip()):
                ref_entries += 1

    print(f"  intro citations      : {intro}")
    print(f"  survey prose         : {survey_prose}")
    print(f"  survey ascending     : {survey_order == sorted(survey_order)}")
    print(f"  table citations      : {table_nums}")
    print(f"  reference entries    : {ref_entries}")
    return problems


def main() -> int:
    ref_path = Path(sys.argv[1] if len(sys.argv) > 1 else REF_DEFAULT)
    gen_path = Path(sys.argv[2] if len(sys.argv) > 2 else GEN_DEFAULT)

    for p in (ref_path, gen_path):
        if not p.exists():
            print(f"missing: {p}")
            return 2

    ref = describe("PRISM (reference)", ref_path)
    gen = describe("SymptoGraph (generated)", gen_path)

    print("=" * 72)
    print(f"{'metric':<24}{'reference':>22}{'generated':>22}")
    print("-" * 72)
    rows = [
        ("page size (in)", ref["page"], gen["page"]),
        ("margins l/t (in)", ref["margins"], gen["margins"]),
        ("paragraphs", ref["paragraphs"], gen["paragraphs"]),
        ("body characters", ref["chars"], gen["chars"]),
        ("tables", ref["tables"], gen["tables"]),
        ("media images", ref["media"], gen["media"]),
        ("OMML equations", ref["omml"], gen["omml"]),
        ("figure captions", len(ref["fig_captions"]), len(gen["fig_captions"])),
        ("table captions", len(ref["tbl_captions"]), len(gen["tbl_captions"])),
        ("numbered headings", len(ref["headings"]), len(gen["headings"])),
        ("X.XXX placeholders", ref["placeholders"], gen["placeholders"]),
    ]
    for name, a, b in rows:
        print(f"{name:<24}{str(a):>22}{str(b):>22}")
    print("=" * 72)

    delta = gen["chars"] - ref["chars"]
    pct = 100.0 * delta / ref["chars"]
    print(f"\nbody size delta: {delta:+d} chars ({pct:+.1f}% vs reference)")

    print("\nreference headings:")
    for num, title in ref["headings"]:
        print(f"  {num:<6} {title}")
    print("\ngenerated headings:")
    for num, title in gen["headings"]:
        print(f"  {num:<6} {title}")

    ref_nums = [n for n, _ in ref["headings"]]
    gen_nums = [n for n, _ in gen["headings"]]
    print(f"\nheading numbering identical: {ref_nums == gen_nums}")
    if ref_nums != gen_nums:
        print(f"  reference only: {sorted(set(ref_nums) - set(gen_nums))}")
        print(f"  generated only: {sorted(set(gen_nums) - set(ref_nums))}")

    print("\ngenerated table shapes (rows x cols):")
    for i, (r, c) in enumerate(gen["table_shapes"], 1):
        print(f"  Table {i:<3} {r:>3} x {c}")
    print("\nreference table shapes (rows x cols):")
    for i, (r, c) in enumerate(ref["table_shapes"], 1):
        print(f"  Table {i:<3} {r:>3} x {c}")

    print("\ncitation audit (generated):")
    problems = audit_citations(gen)

    print("\nfigure captions (generated):")
    for cap in gen["fig_captions"]:
        print(f"  {cap[:88]}")

    if problems:
        print("\nPROBLEMS")
        for p in problems:
            print(f"  - {p}")
        return 1
    print("\nAll citation rules satisfied.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
