#!/usr/bin/env python3
"""Dump per-section prose sizes and table headers for two documents side by side.

Used to drive the manuscript to the same content size and table geometry as the
reference template.  Not part of the build; a diagnostic only.

Usage:  python3 scripts/inspect_ref.py [reference.docx] [generated.docx]
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document

HEADING_RE = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+(.*)$")


def is_heading(text: str) -> bool:
    t = text.strip()
    if not t or len(t) > 90 or t.endswith("."):
        return False
    return bool(HEADING_RE.match(t))


def section_sizes(path: Path) -> list[tuple[str, int, int]]:
    """(section label, prose chars, paragraph count) walking the body in order."""
    doc = Document(str(path))
    out: list[tuple[str, int, int]] = []
    label = "front matter"
    chars = 0
    paras = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if is_heading(text):
            out.append((label, chars, paras))
            label = text
            chars = 0
            paras = 0
            continue
        if text:
            chars += len(text)
            paras += 1
    out.append((label, chars, paras))
    return out


def table_headers(path: Path) -> list[tuple[int, int, list[str]]]:
    doc = Document(str(path))
    rows = []
    for t in doc.tables:
        header = [c.text.strip().replace("\n", " ")[:28] for c in t.rows[0].cells]
        rows.append((len(t.rows), len(t.columns), header))
    return rows


def first_column(path: Path, idx: int) -> list[str]:
    doc = Document(str(path))
    t = doc.tables[idx]
    return [r.cells[0].text.strip().replace("\n", " ")[:46] for r in t.rows]


def main() -> int:
    ref = Path(sys.argv[1] if len(sys.argv) > 1 else "PRISM - Copy.docx")
    gen = Path(sys.argv[2] if len(sys.argv) > 2 else "docs/SymptoGraph-LLM_Paper.docx")

    rs = section_sizes(ref)
    gs = section_sizes(gen)

    print("=" * 94)
    print("PER-SECTION PROSE SIZE (chars, paragraphs)   reference -> generated")
    print("=" * 94)
    gmap = {k: (c, p) for k, c, p in gs}
    # match on heading number so renamed titles still align
    def key(label: str) -> str:
        m = HEADING_RE.match(label.strip())
        return m.group(1) if m else label.strip().lower()

    gkeyed = {key(k): (k, c, p) for k, c, p in gs}
    total_r = total_g = 0
    for label, c, p in rs:
        k = key(label)
        gc, gp = (gkeyed[k][1], gkeyed[k][2]) if k in gkeyed else (0, 0)
        total_r += c
        total_g += gc
        delta = gc - c
        flag = "   <-- short" if delta < -600 else ("   <-- long" if delta > 600 else "")
        print(f"{label[:44]:<46}{c:>7} ({p:>2}p) -> {gc:>7} ({gp:>2}p)  {delta:+7d}{flag}")
    print("-" * 94)
    print(f"{'TOTAL':<46}{total_r:>7}        -> {total_g:>7}        {total_g-total_r:+7d}")

    print("\n" + "=" * 94)
    print("TABLE GEOMETRY AND HEADERS")
    print("=" * 94)
    rt = table_headers(ref)
    gt = table_headers(gen)
    for i in range(max(len(rt), len(gt))):
        print(f"\nTable {i+1}")
        if i < len(rt):
            r, c, h = rt[i]
            print(f"  ref {r:>3}x{c}  {h}")
        if i < len(gt):
            r, c, h = gt[i]
            print(f"  gen {r:>3}x{c}  {h}")

    for idx in (3, 5, 6, 8, 9):
        if idx < len(rt):
            print(f"\n--- reference Table {idx+1} first column ---")
            for v in first_column(ref, idx):
                print(f"    {v}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
