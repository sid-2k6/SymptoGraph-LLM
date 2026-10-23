#!/usr/bin/env python3
"""Audit metric consistency between the prose and the tables.

Every performance number quoted in the body text must also appear in one of the
tables, so a reader cross-checking a sentence against a table never finds two
different values for the same quantity. Numbers that legitimately appear only in
prose - dataset counts, corpus overlaps, configuration constants, measured
diagnostics that are not table rows - are whitelisted explicitly.

Exits non-zero if an unexplained prose number is found.

Usage:  python3 scripts/check_consistency.py [docx]
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document

DEFAULT = "docs/SymptoGraph-LLM_Paper.docx"

# Three-decimal values that are deliberately prose-only, each with a reason.
WHITELIST_EXACT = {
    # graph prior edge weights quoted in 3.4, a property of the prior not a result
    "0.676", "0.428", "0.404",
    # mean distinct conditions per author, justifying the conditional estimator
    "1.125",
    # threshold-policy illustration in 4.2, measured during development
    "0.336", "0.603",
    # contrastive and variance loss weights, configuration not performance
    "0.100", "0.500",
    # margins obtained by subtracting two values that are both in the tables,
    # so a reader can verify them directly from Table 10
    "0.028", "0.079",
}

# patterns whose numbers are structural rather than performance figures
SKIP_LINE = re.compile(
    r"^(Table|Figure|\(\d+\)|\[\d+\])", re.IGNORECASE)

NUM = re.compile(r"(?<![\d.])(\d\.\d{3})(?![\d])")


def table_numbers(doc: Document) -> set[str]:
    out: set[str] = set()
    for tb in doc.tables:
        for row in tb.rows:
            for cell in row.cells:
                for m in NUM.finditer(cell.text):
                    out.add(m.group(1))
    return out


def prose_numbers(doc: Document) -> list[tuple[str, str]]:
    """[(value, containing sentence)] for body paragraphs only."""
    found: list[tuple[str, str]] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text or SKIP_LINE.match(text):
            continue
        for m in NUM.finditer(text):
            start = text.rfind(".", 0, m.start()) + 1
            end = text.find(".", m.end())
            sentence = text[start: end if end != -1 else len(text)].strip()
            found.append((m.group(1), sentence))
    return found


def main() -> int:
    path = Path(sys.argv[1] if len(sys.argv) > 1 else DEFAULT)
    if not path.exists():
        print(f"missing: {path}")
        return 2

    doc = Document(str(path))
    in_tables = table_numbers(doc)
    in_prose = prose_numbers(doc)

    unexplained: list[tuple[str, str]] = []
    for value, sentence in in_prose:
        if value in in_tables or value in WHITELIST_EXACT:
            continue
        unexplained.append((value, sentence))

    distinct_prose = {v for v, _ in in_prose}
    print("=" * 72)
    print("METRIC CONSISTENCY AUDIT")
    print("=" * 72)
    print(f"  distinct 3-decimal values in tables : {len(in_tables)}")
    print(f"  3-decimal values quoted in prose    : {len(in_prose)} "
          f"({len(distinct_prose)} distinct)")
    print(f"  prose values backed by a table      : "
          f"{len(distinct_prose & in_tables)}")
    print(f"  whitelisted prose-only values       : "
          f"{len(distinct_prose & WHITELIST_EXACT)}")
    print(f"  unexplained                         : {len(unexplained)}")

    if unexplained:
        print("\nUNEXPLAINED PROSE NUMBERS")
        for value, sentence in unexplained:
            print(f"  {value}  <-  {sentence[:140]}")
        return 1

    print("\nEvery performance number in the prose is backed by a table.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
