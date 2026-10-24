#!/usr/bin/env python3
"""Confirm equations, figures and tables appear in ascending order in the body,
and count placeholders in prose versus tables.

Usage:  python3 scripts/check_doc_order.py [docx]
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from docx import Document

DEFAULT = "docs/SymptoGraph-LLM_Paper.docx"


def main() -> int:
    path = Path(sys.argv[1] if len(sys.argv) > 1 else DEFAULT)
    doc = Document(str(path))

    eq_nums: list[int] = []
    fig_nums: list[int] = []
    tbl_nums: list[int] = []
    prose_ph = 0

    for para in doc.paragraphs:
        t = para.text.strip()
        prose_ph += t.count("X.XXX")
        if re.fullmatch(r"\((\d{1,2})\)", t):
            eq_nums.append(int(t[1:-1]))
            continue
        m = re.match(r"^Figure\s+(\d+)\.", t)
        if m:
            fig_nums.append(int(m.group(1)))
        m = re.match(r"^Table\s+(\d+)\.", t)
        if m:
            tbl_nums.append(int(m.group(1)))

    table_ph = 0
    for tb in doc.tables:
        for row in tb.rows:
            for cell in row.cells:
                table_ph += cell.text.count("X.XXX")

    with zipfile.ZipFile(path) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]

    print(f"equation numbers in body : {eq_nums}")
    print(f"  ascending 1..20        : {eq_nums == list(range(1, 21))}")
    print(f"figure captions          : {fig_nums}")
    print(f"  ascending 1..5         : {fig_nums == list(range(1, 6))}")
    print(f"table captions           : {tbl_nums}")
    print(f"  ascending 1..10        : {tbl_nums == list(range(1, 11))}")
    print(f"embedded media files     : {len(media)}")
    print(f"placeholders in prose    : {prose_ph}")
    print(f"placeholders in tables   : {table_ph}")
    print(f"placeholders total       : {prose_ph + table_ph}")

    with zipfile.ZipFile(path) as zf:
        doc_xml = zf.read("word/document.xml").decode("utf-8", "replace")
    n_omml = doc_xml.count("<m:oMath")

    print(f"native OMML equations    : {n_omml}")

    # 7 media files: figures 1, 2, 3a, 3b, 4, 5a, 5b. Equations are OMML now,
    # so an equation image appearing here would mean a regression.
    ok = (
        eq_nums == list(range(1, 21))
        and fig_nums == list(range(1, 6))
        and tbl_nums == list(range(1, 11))
        and len(media) == 7
        and n_omml == 20
        and prose_ph + table_ph == 0
    )
    print(f"\nall ordering checks pass : {ok}")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
